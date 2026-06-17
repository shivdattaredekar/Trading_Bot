import json
import os
from datetime import datetime
from typing import Dict, List

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from tradingbot.utils.logger import log

# --------------------------------------------------
# CONFIG
# --------------------------------------------------
ACTIVE_TRADES_FILE = "active_trades.json"
OUTPUT_FILE = "pnl_tracker.xlsx"
SHEET_NAME = "pnl_tracker"

FYERS_TIME_FORMAT = "%d-%b-%Y %H:%M:%S"
OUTPUT_TIME_FORMAT = "%Y-%m-%d %H.%M.%S"



# --------------------------------------------------
# LOAD ACTIVE TRADES (JSON)
# --------------------------------------------------
def load_active_trades(path: str = ACTIVE_TRADES_FILE) -> Dict[str, dict]:
    try:
        with open(path, "r") as f:
            data = json.load(f)
        log(f"📥 Loaded active_trades from {path}")
        return data
    except FileNotFoundError:
        log(f"❌ active_trades.json not found at {path}")
        return {}
    except json.JSONDecodeError as e:
        log(f"❌ Invalid JSON in active_trades.json: {e}")
        return {}


# --------------------------------------------------
# FETCH TRADEBOOK FROM FYERS
# --------------------------------------------------
def fetch_tradebook(fyers) -> List[dict]:

    resp = fyers.tradebook()
    if resp.get("code") != 200:
        log(f"❌ Failed to fetch tradebook: {resp}")
        return []
    return resp.get("tradeBook", [])


# --------------------------------------------------
# BUILD TRADE-WISE PNL (ACTIVE_TRADES DRIVEN)
# --------------------------------------------------
def build_tradewise_pnl(
    tradebook: List[dict],
    active_trades: Dict[str, dict],
) -> List[Dict]:

    results = []

    for order_id, trade in active_trades.items():
        if trade.get("status") != "CLOSED":
            continue

        symbol = trade["symbol"]
        side = int(trade["side"])           # 1 = LONG, -1 = SHORT
        qty = int(trade["qty"])
        entry_price = float(trade["entry_price"])
        stop_loss = float(trade["stop_price"])

        created_at = datetime.fromisoformat(trade["created_at"])
        closed_at = datetime.fromisoformat(trade["closed_at"])

        exit_trades = [
            t for t in tradebook
            if t["symbol"] == symbol
            and int(t["side"]) == -side
            and created_at
            <= datetime.strptime(t["orderDateTime"], FYERS_TIME_FORMAT)
            <= closed_at
        ]

        if not exit_trades:
            log(f"⚠️ No exit trades found for order {order_id}")
            continue

        total_exit_qty = sum(int(t["tradedQty"]) for t in exit_trades)
        total_exit_value = sum(
            int(t["tradedQty"]) * float(t["tradePrice"])
            for t in exit_trades
        )

        avg_exit_price = total_exit_value / total_exit_qty

        pnl = (
            (avg_exit_price - entry_price) * qty
            if side == 1
            else (entry_price - avg_exit_price) * qty
        )

        results.append({
            "timestamp": closed_at.strftime(OUTPUT_TIME_FORMAT),
            "symbol": symbol,
            "qty": qty,
            "entry_price": round(entry_price, 2),
            "stop_loss": round(stop_loss, 2),
            "target": None,
            "PnL": round(pnl, 2),
        })

    return results


# --------------------------------------------------
# APPEND TO EXCEL (SAFE, DAILY RUN)
# --------------------------------------------------
def export_to_excel(rows: List[Dict]):
    headers = ["timestamp", "symbol", "qty", "entry_price", "stop_loss", "target", "PnL"]

    if os.path.exists(OUTPUT_FILE):
        wb = load_workbook(OUTPUT_FILE)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.create_sheet(SHEET_NAME)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        ws.append(headers)

    for r in rows:
        ws.append([
            r["timestamp"],
            r["symbol"],
            r["qty"],
            r["entry_price"],
            r["stop_loss"],
            r["target"],
            r["PnL"],
        ])

    # Auto-adjust column width
    for col in ws.columns:
        max_len = max(len(str(c.value)) if c.value else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 2

    wb.save(OUTPUT_FILE)
    log(f"📤 Appended {len(rows)} trades to {OUTPUT_FILE}")

